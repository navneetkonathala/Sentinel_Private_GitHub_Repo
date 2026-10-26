"""Build Sentinel's official 10-slide Utkrishta Challenge 2 presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Sentinel_Challenge_2_Presentation.pptx"
SW, SH = 13.333, 7.5
INK, DEEP, TEAL = "172B2A", "102320", "168A75"
MINT, PAPER, WHITE = "DDF1E9", "F5F4EF", "FFFFFF"
MUTED, LINE, ORANGE = "718079", "DDE2DA", "D9794E"
GOLD, PALE_BLUE, PALE_ORANGE, PALE_YELLOW = "E7BC60", "EAF0EC", "F9E9E0", "F6F0DF"
FONT, DISPLAY, MONO = "Aptos", "Aptos Display", "Cascadia Code"


def color(hex_value):
    return RGBColor.from_string(hex_value)


def shape(slide, x, y, w, h, fill, line=None, rounded=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    item = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    item.fill.solid()
    item.fill.fore_color.rgb = color(fill)
    item.line.color.rgb = color(line or fill)
    if rounded:
        item.adjustments[0] = .06
    return item


def text(slide, value, x, y, w, h, size=12, ink=INK, bold=False, font=FONT,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color(ink)
    return box


def line(slide, x1, y1, x2, y2, ink=LINE, width=1.0):
    item = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    item.line.color.rgb = color(ink)
    item.line.width = Pt(width)
    return item


def bullet(slide, label, x, y, w, size=11, ink=INK, dot=TEAL, h=.45):
    shape(slide, x, y + .07, .07, .07, dot, dot, True)
    text(slide, label, x + .18, y, w - .18, h, size, ink)


def pill(slide, label, x, y, w, fill=MINT, ink=TEAL):
    shape(slide, x, y, w, .29, fill, fill, True)
    text(slide, label, x + .07, y + .055, w - .14, .16, 7.5, ink, True,
         MONO, PP_ALIGN.CENTER)


def base(prs, number, kicker, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(PAPER)
    text(slide, "SENTINEL  /  UTKRISHTA 2026", .62, .30, 4.8, .2, 8, MUTED, True, MONO)
    text(slide, kicker.upper(), 9.2, .30, 3.5, .2, 8, TEAL, True, MONO, PP_ALIGN.RIGHT)
    line(slide, .62, .62, 12.70, .62, LINE, .7)
    text(slide, title, .72, .88, 11.9, .56, 26, INK, True, DISPLAY)
    if subtitle:
        text(slide, subtitle, .75, 1.52, 11.7, .45, 11, MUTED)
    line(slide, .62, 7.12, 12.70, 7.12, LINE, .7)
    text(slide, "CHALLENGE 02 · ADAPTIVE TIME-SERIES ANOMALY DETECTION",
         .72, 7.22, 8.2, .14, 7, MUTED, False, MONO)
    text(slide, f"{number:02d} / 10", 11.7, 7.22, 1.0, .14, 7, MUTED,
         False, MONO, PP_ALIGN.RIGHT)
    return slide


def add_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(DEEP)
    text(slide, "UTKRISHTA 2026  /  AI·ML HACKATHON", .76, .50, 5.4, .22, 9,
         "A9C5BA", True, MONO)
    line(slide, .76, .86, 12.55, .86, "36504A", .8)
    pill(slide, "CHALLENGE 02", .78, 1.32, 1.48, "27453E", "83D7B8")
    text(slide, "Sentinel", .78, 1.94, 6.0, .83, 50, WHITE, True, DISPLAY)
    text(slide, "Learn the usual.\nExplain the unusual.", .82, 2.98, 6.0, 1.1,
         27, "B8D0C5", False, DISPLAY)
    text(slide, "A lightweight, streaming anomaly detector using a permitted\nsliding-window adaptive threshold.",
         .84, 4.43, 5.7, .67, 14, "A8BAB1")
    shape(slide, .82, 5.63, 4.25, .73, "223B36", "38544C", True)
    text(slide, "TEAM NAME: ________________________", 1.02, 5.83, 3.9, .16,
         9, "C5D8CE", True, MONO)
    text(slide, "TEAM MEMBERS: __________________________________________", .83,
         6.62, 7.3, .18, 8.5, "A8BAB1", False, MONO)
    # Signal visual: normal range with one high anomaly.
    cx, cy, cw, ch = 7.55, 2.15, 4.45, 2.92
    for i in range(5):
        yy = cy + i * ch / 4
        line(slide, cx, yy, cx + cw, yy, "2E4842", .6)
    line(slide, cx, cy + 1.43, cx + cw, cy + 1.43, "79958A", 1.2)
    line(slide, cx, cy + .70, cx + cw, cy + .70, "45635A", .8)
    line(slide, cx, cy + 2.16, cx + cw, cy + 2.16, "45635A", .8)
    values = [.48,.45,.50,.43,.47,.41,.45,.50,.44,.49,.42,.46,.44,.48,.43,.47,.45,.10,.46,.42,.47]
    pts = [(cx + i*cw/(len(values)-1), cy + .48 + v*2.0) for i,v in enumerate(values)]
    for a,b in zip(pts,pts[1:]):
        line(slide,a[0],a[1],b[0],b[1],"55D0AA",2)
    anomaly = pts[17]
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(anomaly[0]-.075), Inches(anomaly[1]-.075), Inches(.15), Inches(.15))
    dot.fill.solid(); dot.fill.fore_color.rgb=color(ORANGE); dot.line.color.rgb=color(ORANGE)
    text(slide, "SIMULATED SIGNAL  ·  ONE FLAGGED POINT", 7.58, 5.36, 4.7, .19,
         8, "A8C5B8", True, MONO)
    text(slide, "PRESENTATION DRAFT  ·  REPLACE TEAM DETAILS BEFORE SUBMISSION", .78,
         7.16, 8.5, .16, 7, "82958C", False, MONO)


def slide_problem(prs):
    s = base(prs, 2, "Problem understanding", "What does Challenge 2 ask us to solve?",
             "Watch measurements arrive in time order, learn what is usual, and explain readings or periods that differ significantly.")
    shape(s,.75,2.15,7.25,2.15,WHITE,LINE,True)
    pill(s,"THE INPUT",1.02,2.40,1.0)
    text(s,"A time series",1.03,2.87,2.5,.35,19,INK,True,DISPLAY)
    text(s,"Ordered readings, each with a time and numeric value.",1.04,3.34,2.6,.54,11,MUTED)
    text(s,"→",3.72,3.08,.55,.4,25,TEAL,True,DISPLAY,PP_ALIGN.CENTER)
    text(s,"An anomaly",4.48,2.87,2.6,.35,19,INK,True,DISPLAY)
    text(s,"A point or time window far from the learned normal pattern.",4.49,3.34,2.8,.54,11,MUTED)
    shape(s,8.25,2.15,4.3,2.15,INK,INK,True)
    text(s,"EACH FLAG NEEDS",8.55,2.45,3.5,.19,9,"9BC4B6",True,MONO)
    bullet(s,"Anomaly score",8.58,2.87,3.5,11,WHITE,"47C5A0")
    bullet(s,"NORMAL / ANOMALOUS label",8.58,3.27,3.5,11,WHITE,"47C5A0")
    bullet(s,"Evidence a person can interpret",8.58,3.67,3.5,11,WHITE,"47C5A0")
    text(s,"Why it matters",.84,4.70,3.0,.3,17,INK,True,DISPLAY)
    cards=[("INDUSTRY","Spot unusual machine or process readings."),("CLOUD","Notice a sudden jump in latency or memory."),("SENSORS","Find faulty measurements without alerting on every wiggle.")]
    for i,(head,body) in enumerate(cards):
        x=.82+i*4.0
        shape(s,x,5.18,3.72,1.05,[MINT,"F6F0DF",PALE_ORANGE][i],LINE,True)
        text(s,head,x+.17,5.36,3.3,.18,8,TEAL if i<2 else ORANGE,True,MONO)
        text(s,body,x+.17,5.66,3.35,.4,10.5,INK)
    text(s,"Sentinel's current project is a prototype; no real company feed is connected.",.85,6.52,11.7,.2,9,MUTED,False,FONT,italic=True)


def slide_data(prs):
    s=base(prs,3,"Dataset & preparation","The benchmark is a labeled timeline—not random demo data",
            "The brief names a prepared Numenta Anomaly Benchmark (NAB) subset with hidden test windows.")
    shape(s,.78,2.12,5.2,3.85,WHITE,LINE,True)
    text(s,"ILLUSTRATIVE ROW SHAPE",1.05,2.4,4.5,.2,9,TEAL,True,MONO)
    headers=[("TIMESTAMP",1.08),("VALUE",2.82),("LABEL*",4.4)]
    for label,x in headers:text(s,label,x,2.83,1.3,.17,8,MUTED,True,MONO)
    rows=[("10:00:01","40.1","normal"),("10:00:02","41.2","normal"),("10:00:03","39.8","normal"),("10:00:04","58.0","event window"),("10:00:05","40.5","normal")]
    for i,row in enumerate(rows):
        y=3.20+i*.45
        line(s,1.05,y+.27,5.65,y+.27,"EDF0EB",.6)
        text(s,row[0],1.08,y,1.4,.18,9.5,INK,False,MONO)
        text(s,row[1],2.84,y,.8,.18,9.5,INK,False,MONO)
        text(s,row[2],4.4,y,1.2,.18,8.5,ORANGE if i==3 else TEAL,True,MONO)
    text(s,"*Example only. Labels refer to organizer-supplied anomaly windows; no NAB file is in this workspace.",1.05,5.58,4.6,.28,8.5,MUTED,False,FONT,italic=True)
    text(s,"Preparation workflow",6.55,2.28,5.3,.3,18,INK,True,DISPLAY)
    flow=[("01","LOAD","Read timestamp, value, event label."),("02","CHECK","Inspect gaps, duplicates, units, ordering."),("03","STREAM","Feed past and current values only."),("04","SCORE","Compare alerts with labeled windows.")]
    for i,(num,head,body) in enumerate(flow):
        y=2.92+i*.77
        pill(s,num,6.58,y,.48,MINT,TEAL)
        text(s,head,7.27,y+.03,1.1,.18,8.5,TEAL,True,MONO)
        text(s,body,8.42,y,3.75,.3,10,MUTED)
    shape(s,6.56,6.1,5.7,.42,PALE_ORANGE,PALE_ORANGE,True)
    text(s,"NEXT REQUIRED INPUT: organizer's NAB subset + schema",6.75,6.22,5.3,.15,8,ORANGE,True,MONO)


def slide_system(prs):
    s=base(prs,4,"System design","From one reading to an explained event",
            "The Python detector is the reference algorithm; the browser runs a separate synthetic demo using matching settings.")
    cards=[("01","INPUT","timestamp\n+ numeric value"),("02","WINDOW","51 recent\nclean readings"),("03","SCORE","median/MAD\nrobust score"),("04","DECIDE","NORMAL or\nANOMALOUS"),("05","EXPLAIN","limits, score,\ndirection, reason")]
    widths=[2.0,2.0,2.0,2.0,2.0]; start=.72; gap=.45
    for i,(n,h,b) in enumerate(cards):
        x=start+i*(widths[i]+gap)
        shape(s,x,2.55,widths[i],1.62,WHITE,LINE,True)
        pill(s,n,x+.16,2.73,.45,MINT,TEAL)
        text(s,h,x+.16,3.12,widths[i]-.32,.2,9,TEAL,True,MONO)
        text(s,b,x+.16,3.48,widths[i]-.3,.48,11,INK,True,DISPLAY)
        if i<4:text(s,"→",x+widths[i]+.06,3.14,.34,.28,17,TEAL,True,DISPLAY,PP_ALIGN.CENTER)
    shape(s,.78,4.75,5.75,1.12,"E6F2ED","D5E7DE",True)
    text(s,"PYTHON REFERENCE",1.03,4.98,2.3,.18,8,TEAL,True,MONO)
    text(s,"sentinel.py  →  Sentinel.update(value)  →  result dictionary",1.03,5.32,5.05,.24,10,INK,False,MONO)
    shape(s,6.78,4.75,5.75,1.12,INK,INK,True)
    text(s,"BROWSER DEMO",7.03,4.98,2.3,.18,8,"8ED6BC",True,MONO)
    text(s,"app.js  →  generated signal  →  chart + event log",7.03,5.32,5.1,.24,10,WHITE,False,MONO)
    text(s,"The website does not ingest NAB or call the Python class yet.",.84,6.25,11.5,.22,10,ORANGE,True)


def slide_algorithm(prs):
    s=base(prs,5,"Core algorithm","A permitted sliding-window adaptive threshold",
           "The brief explicitly permits sliding-window adaptive thresholds. Sentinel uses median + MAD for a robust, interpretable limit.")
    shape(s,.78,2.15,5.2,3.55,INK,INK,True)
    text(s,"ROBUST SCORE",1.07,2.48,3.4,.19,9,"8ED6BC",True,MONO)
    text(s,"center = median(window)\nMAD = median(|x − center|)\nspread = max(1.4826 × MAD, ε)\nscore = (new − center) / spread",
         1.08,2.94,4.65,1.45,14,WHITE,False,MONO)
    text(s,"Flag if |score| > 3.5",1.08,4.72,4.45,.28,15,"F2C089",True,MONO)
    text(s,"CURRENT SETTINGS",6.58,2.35,4.0,.18,9,TEAL,True,MONO)
    settings=[("Window","51 accepted readings"),("Warm-up","12 readings"),("Threshold","3.5 robust-score units"),("Update rule","Do not learn from flagged points")]
    for i,(a,b) in enumerate(settings):
        y=2.82+i*.62
        text(s,a,6.62,y,1.55,.19,10,INK,True)
        text(s,b,8.25,y,3.7,.23,10,MUTED)
        line(s,6.6,y+.35,12.14,y+.35,LINE,.6)
    shape(s,6.58,5.48,5.56,.52,MINT,MINT,True)
    text(s,"WHY THIS FITS: low compute · adapts online · explains flags",6.77,5.66,5.2,.16,8,TEAL,True,MONO)
    text(s,"Fit-based choice, not an evidence-backed winner: NAB comparison is still pending.",.84,6.38,11.4,.22,9,ORANGE,False,FONT,italic=True)


def slide_experiments(prs):
    s=base(prs,6,"Tuning plan","Honest experiment log: no NAB runs have been completed",
           "The browser spike is an interaction demo, not a benchmark run. We will not present it as predictive performance.")
    shape(s,.78,2.14,7.55,3.88,WHITE,LINE,True)
    heads=[("CANDIDATE",1.03,2.4),("SETTING TO TRY",4.12,2.2),("STATUS",6.66,1.3)]
    for h,x,w in heads:text(s,h,x,2.45,w,.18,8,MUTED,True,MONO)
    rows=[("Rolling z-score","mean/std; 51-point window","Planned"),("Sentinel robust threshold","median/MAD; 51, cutoff 3.5","Current prototype"),("Threshold sweep","3.0 / 3.5 / 4.0 on validation","Planned"),("Window sweep","25 / 51 / 101 clean points","Planned")]
    for i,(a,b,c) in enumerate(rows):
        y=2.91+i*.64
        line(s,1.03,y-.12,8.02,y-.12,"EDF0EB",.6)
        text(s,a,1.04,y,2.9,.27,9.5,INK,True)
        text(s,b,4.12,y,2.35,.34,8.5,MUTED)
        text(s,c,6.66,y,1.3,.3,8,ORANGE if c!="Current prototype" else TEAL,True)
    shape(s,8.66,2.14,3.88,3.88,INK,INK,True)
    text(s,"NEXT STEPS",8.97,2.48,3.0,.18,9,"8ED6BC",True,MONO)
    bullet(s,"Obtain the prepared NAB subset.",8.98,2.98,3.12,10,WHITE,"47C5A0",.52)
    bullet(s,"Keep time order; tune on validation only.",8.98,3.72,3.12,10,WHITE,"47C5A0",.58)
    bullet(s,"Log settings, runtime, and every result.",8.98,4.55,3.12,10,WHITE,"47C5A0",.55)
    text(s,"Do not tune on hidden test windows.",8.98,5.45,3.1,.22,9,"F2C089",True)
    text(s,"Challenge 2 has no 15-run cap in the supplied brief; that limit belongs to Challenge 1.",.85,6.34,11.5,.24,9,MUTED,False,FONT,italic=True)


def slide_metrics(prs):
    s=base(prs,7,"Evaluation methodology","Measure events, delay, and cost—not just a dramatic spike",
           "The brief names event-based F1 as primary, with precision, recall, detection delay, and inference runtime.")
    metrics=[("PRECISION","Of alerts, how many match real events?","TP / (TP + FP)",PALE_ORANGE),
             ("RECALL","Of real events, how many are detected?","TP / (TP + FN)",MINT),
             ("EVENT F1","Balance precision and recall across windows.","2PR / (P + R)",PALE_BLUE),
             ("DETECTION DELAY","Time from event onset to first alert.","alert time − onset",PALE_YELLOW)]
    for i,(h,body,formula,fill) in enumerate(metrics):
        x=.82+(i%2)*6.0; y=2.22+(i//2)*1.68
        shape(s,x,y,5.65,1.38,fill,fill,True)
        text(s,h,x+.22,y+.2,4.8,.18,8,TEAL if i!=0 else ORANGE,True,MONO)
        text(s,body,x+.22,y+.52,5.1,.28,12,INK,True,DISPLAY)
        text(s,formula,x+.22,y+.94,4.9,.2,9,MUTED,False,MONO)
    shape(s,.82,5.82,11.65,.56,INK,INK,True)
    text(s,"RUNTIME",1.08,6.02,1.0,.17,8,"8ED6BC",True,MONO)
    text(s,"Report per-sample inference time on the same hardware and input windows.",2.25,5.99,9.7,.22,10,WHITE)
    text(s,"Choose event-matching rules before evaluation and preserve chronological splits.",.86,6.61,11.5,.18,9,MUTED,False,FONT,italic=True)


def slide_results(prs):
    s=base(prs,8,"Results & comparison","Results are pending the official NAB data",
           "No hidden-test or NAB validation scores have been produced in this workspace. The table is intentionally not filled with demo numbers.")
    shape(s,.82,2.14,7.25,3.82,WHITE,LINE,True)
    cols=[("METHOD",1.08,2.45),("EVENT F1",4.0,1.15),("RUNTIME",5.43,1.2),("STATUS",6.76,1.0)]
    for h,x,w in cols:text(s,h,x,2.48,w,.18,8,MUTED,True,MONO)
    methods=[("Rolling mean/std z-score","—","—","Not run"),("Sentinel: median + MAD","—","—","Prototype only"),("Isolation Forest / LOF*","—","—","Optional baseline")]
    for i,row in enumerate(methods):
        y=3.02+i*.75
        line(s,1.06,y-.12,7.84,y-.12,"EDF0EB",.6)
        for j,(val,(_,x,w)) in enumerate(zip(row,cols)):
            text(s,val,x,y,w,.35,9,TEAL if j==3 and i==1 else INK, j==0 or j==3, MONO if j in (1,2,3) else FONT)
    text(s,"*Run only if the chosen NAB subset and runtime budget make them appropriate.",1.08,5.48,6.5,.2,8,MUTED,False,FONT,italic=True)
    shape(s,8.43,2.14,4.03,3.82,INK,INK,True)
    text(s,"WHAT WE CAN SHOW TODAY",8.75,2.5,3.3,.18,8,"8ED6BC",True,MONO)
    text(s,"A synthetic reading is flagged and the UI explains why.",8.75,3.0,3.25,.62,15,WHITE,True,DISPLAY)
    text(s,"This demonstrates interaction and evidence formatting—not accuracy on real events.",8.75,4.06,3.15,.72,10,"B8C9C0")
    pill(s,"NO PERFORMANCE CLAIMS",8.75,5.25,2.35,"3A332B","F2C089")
    text(s,"Before submission: populate this slide from reproducible validation and organizer results.",.86,6.35,11.5,.2,9,ORANGE,True)


def slide_demo_access(prs):
    s=base(prs,9,"Demo & anomaly evidence","A live demo with honest data labels and role-based access",
           "The website currently generates synthetic values. No company database, server packets, or live sensor stream is connected.")
    # mini dashboard mock
    shape(s,.76,2.12,6.35,3.95,WHITE,LINE,True)
    text(s,"SENTINEL  /  SAMPLE EVENT",1.0,2.39,4.9,.18,8,TEAL,True,MONO)
    text(s,"58.0",1.02,2.82,1.4,.42,25,INK,True,DISPLAY)
    text(s,"observed value",1.05,3.28,1.5,.18,8,MUTED)
    text(s,"Baseline 40.55   ·   Expected range 38.21–42.89",2.65,2.93,3.95,.3,10,INK,True)
    shape(s,2.66,3.45,3.85,.43,PALE_ORANGE,PALE_ORANGE,True)
    text(s,"HIGH ANOMALY  ·  +15.1 ABOVE UPPER LIMIT",2.82,3.59,3.5,.15,8,ORANGE,True,MONO)
    text(s,"Robust score: +26.16",1.05,4.12,2.9,.22,10,TEAL,True,MONO)
    text(s,"Illustrative output from the built-in example; not a NAB result.",1.05,4.42,5.5,.18,8,MUTED,False,FONT,italic=True)
    # role flow
    text(s,"HOW PEOPLE USE THE SITE",7.55,2.26,4.65,.2,9,TEAL,True,MONO)
    shape(s,7.55,2.68,4.9,1.15,MINT,"D4E7DE",True)
    text(s,"PRESENTER",7.79,2.91,1.4,.18,8,TEAL,True,MONO)
    text(s,"Sign in → add presenters → open member access",7.79,3.23,4.28,.27,10,INK,True)
    shape(s,7.55,4.08,4.9,1.15,PALE_BLUE,"D9E3DD",True)
    text(s,"INSTITUTE MEMBER",7.79,4.31,1.8,.18,8,TEAL,True,MONO)
    text(s,"Register @iiits.in → sign in → use demo controls after access opens",7.79,4.62,4.3,.42,9.5,INK,True)
    text(s,"The email domain is checked; mailbox ownership is not verified.",7.57,5.47,4.8,.23,8,ORANGE,False,FONT,italic=True)
    text(s,"Demo controls: inject anomaly · simulate drift · pause · reset. Members see them only after the presenter opens access.",.85,6.4,11.5,.3,9,MUTED)


def slide_limits_run(prs):
    s=base(prs,10,"Run, limitations & next steps","Run the program locally—and know what it does not prove",
           "The Python server protects dashboard routes; do not use a static file server for authenticated access.")
    shape(s,.78,2.12,5.55,3.96,INK,INK,True)
    text(s,"LOCAL DEMO",1.08,2.43,3.0,.18,9,"8ED6BC",True,MONO)
    code="1  python3 auth_server.py\n2  open http://127.0.0.1:8001\n\nSign in with the presenter\ncredentials in AUTH_SETUP.md\n\nDetector only: python3 sentinel.py"
    text(s,code,1.08,2.9,4.8,1.95,11.5,WHITE,False,MONO)
    text(s,"Python 3.9+ · runtime uses the standard library · SQLite DB is created on first server start.",1.08,5.3,4.7,.48,9,"B8C9C0")
    text(s,"NEXT WORK",6.78,2.38,4.3,.18,9,TEAL,True,MONO)
    bullet(s,"Load and inspect the organizer's NAB subset.",6.8,2.85,5.1,10.5,INK,TEAL,.46)
    bullet(s,"Evaluate event-F1, precision, recall, delay, runtime.",6.8,3.53,5.1,10.5,INK,TEAL,.5)
    bullet(s,"Connect a real timestamped stream; validate units and gaps.",6.8,4.25,5.1,10.5,INK,TEAL,.55)
    bullet(s,"Replace the demo password and deploy behind HTTPS.",6.8,5.02,5.1,10.5,INK,TEAL,.5)
    shape(s,6.78,5.83,5.54,.43,PALE_ORANGE,PALE_ORANGE,True)
    text(s,"LIMIT: Python detector and JS demo are separate implementations.",6.96,5.96,5.18,.16,8,ORANGE,True,MONO)
    text(s,"Team: fill in team details on slide 1; use only measured results in any final claims.",.86,6.5,11.3,.22,9,MUTED,False,FONT,italic=True)


def build_deck():
    prs=Presentation()
    prs.slide_width=Inches(SW)
    prs.slide_height=Inches(SH)
    prs.core_properties.title="Sentinel — Utkrishta 2026 Challenge 2"
    prs.core_properties.subject="10-slide adaptive time-series anomaly detection project presentation"
    prs.core_properties.author="Utkrishta 2026 participant team"
    prs.core_properties.keywords="Sentinel, time series, anomaly detection, NAB, sliding-window adaptive threshold"
    add_cover(prs)
    slide_problem(prs)
    slide_data(prs)
    slide_system(prs)
    slide_algorithm(prs)
    slide_experiments(prs)
    slide_metrics(prs)
    slide_results(prs)
    slide_demo_access(prs)
    slide_limits_run(prs)
    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_deck()
